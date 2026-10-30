from datetime import timedelta
from types import SimpleNamespace
from unittest import mock

from django.test import SimpleTestCase, TestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from accounts.models import User
from analytics import classify as cl
from analytics import reports, services
from analytics.models import Event, Visit, Visitor

UA_PHONE = "Mozilla/5.0 (iPhone; CPU iPhone OS 17_0 like Mac OS X) AppleWebKit/605.1.15 Mobile/15E148 Safari/604.1"
UA_DESK = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/120.0 Safari/537.36"
UA_IG = UA_PHONE + " Instagram 300.0"


class ClassifyTests(SimpleTestCase):
    def test_utm_wins_over_referrer_and_is_normalized(self):
        r = cl.classify({"utm_source": "LinkedIn ", "utm_medium": "Paid", "utm_campaign": "Oct-GIS"}, "google.com")
        self.assertEqual((r["source"], r["medium"], r["campaign"]), ("linkedin", "paid", "oct-gis"))

    def test_click_id_without_utm(self):
        self.assertEqual(cl.classify({"gclid": "x"}, "")["source"], "google")
        r = cl.classify({"li_fat_id": "x"}, "")
        self.assertEqual((r["source"], r["medium"], r["click_id"]), ("linkedin", "paid", "li_fat_id"))

    def test_referrers(self):
        cases = {"www.google.com": ("google", "organic"), "google.co.uk": ("google", "organic"),
                 "lnkd.in": ("linkedin", "social"), "l.instagram.com": ("instagram", "social"),
                 "t.me": ("telegram", "social"), "notgoogle.com": ("notgoogle.com", "referral"),
                 "blog.example.org": ("blog.example.org", "referral")}
        for host, expected in cases.items():
            r = cl.classify({}, host)
            self.assertEqual((r["source"], r["medium"]), expected, host)

    def test_direct_and_in_app(self):
        self.assertEqual(cl.classify({}, "")["source"], "direct")
        self.assertEqual(cl.classify({}, "", "instagram")["source"], "instagram")

    def test_bots_and_devices(self):
        self.assertTrue(cl.is_bot("Googlebot/2.1"))
        self.assertTrue(cl.is_bot("TelegramBot (like TwitterBot)"))
        self.assertTrue(cl.is_bot(""))
        self.assertFalse(cl.is_bot(UA_PHONE))
        self.assertEqual(cl.device_info(UA_PHONE)[0], "mobile")
        self.assertEqual(cl.device_info(UA_DESK)[:2], ("desktop", "chrome"))
        self.assertEqual(cl.device_info(UA_IG)[2], "instagram")

    def test_referrer_host_ignores_own_site(self):
        own = {"postgis.ir"}
        self.assertEqual(cl.referrer_host("https://www.postgis.ir/help/", own), "")
        self.assertEqual(cl.referrer_host("https://app.postgis.ir/x", own), "")
        self.assertEqual(cl.referrer_host("https://www.linkedin.com/feed/", own), "linkedin.com")


class MiddlewareTests(TestCase):
    def get(self, url="/", ua=UA_DESK, **extra):
        return self.client.get(url, HTTP_USER_AGENT=ua, **extra)

    def test_first_touch_saved_with_utm(self):
        r = self.get("/?utm_source=LinkedIn&utm_medium=paid&utm_campaign=oct")
        self.assertEqual(r.status_code, 200)
        self.assertIn("gp_vid", r.cookies)
        v = Visitor.objects.get()
        self.assertEqual((v.first_source, v.first_medium, v.first_campaign, v.first_landing),
                         ("linkedin", "paid", "oct", "/"))
        visit = Visit.objects.get()
        self.assertEqual((visit.source, visit.landing_path, visit.device), ("linkedin", "/", "desktop"))
        self.assertEqual(len(visit.ip_hash), 16)
        self.assertNotIn("127.0.0.1", visit.ip_hash)

    def test_same_session_counts_pages_and_keeps_first_touch(self):
        self.get("/?utm_source=telegram&utm_medium=social")
        self.get("/terms/")
        self.get("/help/")
        visit = Visit.objects.get()
        self.assertEqual((visit.pages, visit.last_path), (3, "/help/"))
        self.assertEqual(Visitor.objects.count(), 1)
        self.assertEqual(Visitor.objects.get().visits_count, 1)

    def test_new_visit_after_gap_keeps_first_touch(self):
        self.get("/?utm_source=telegram&utm_medium=social")
        Visit.objects.update(last_seen=timezone.now() - timedelta(hours=1))
        self.get("/")
        self.assertEqual(Visit.objects.count(), 2)
        v = Visitor.objects.get()
        self.assertEqual((v.first_source, v.visits_count), ("telegram", 2))
        self.assertEqual(Visit.objects.order_by("-started_at").first().source, "direct")

    def test_new_campaign_click_starts_new_visit(self):
        self.get("/?utm_source=telegram&utm_medium=social")
        self.get("/?utm_source=linkedin&utm_medium=paid")
        self.assertEqual(Visit.objects.count(), 2)
        self.assertEqual(Visitor.objects.get().first_source, "telegram")

    def test_external_referrer_and_in_app(self):
        self.get("/", HTTP_REFERER="https://www.linkedin.com/feed/")
        self.assertEqual(Visit.objects.get().source, "linkedin")
        self.client.cookies.clear()
        self.get("/", ua=UA_IG)
        self.assertEqual(Visit.objects.order_by("-id").first().in_app, "instagram")

    def test_skips_bots_dnt_private_paths_and_non_html(self):
        self.get(ua="Googlebot/2.1")
        self.get(HTTP_DNT="1")
        self.get("/healthz")
        self.get("/robots.txt")
        self.get("/manage/login/")
        self.assertEqual(Visitor.objects.count(), 0)

    def test_disabled_setting(self):
        with override_settings(ANALYTICS_ENABLED=False):
            self.get()
        self.assertEqual(Visitor.objects.count(), 0)

    def test_tracking_failure_never_breaks_page(self):
        with mock.patch("analytics.services._track", side_effect=RuntimeError("boom")):
            self.assertEqual(self.get().status_code, 200)

    def test_staff_not_tracked(self):
        admin = User.objects.create_admin("root", "Str0ng-Admin-Pass")
        self.client.force_login(admin)
        self.get()
        self.assertEqual(Visitor.objects.count(), 0)

    def test_invalid_cookie_gets_replaced(self):
        self.client.cookies["gp_vid"] = "not-a-real-id"
        r = self.get()
        self.assertEqual(len(r.cookies["gp_vid"].value), 32)


class FunnelTests(TestCase):
    def test_signup_links_visitor_and_records_funnel(self):
        self.client.get("/?utm_source=linkedin&utm_medium=paid&utm_campaign=oct", HTTP_USER_AGENT=UA_DESK)
        with mock.patch("accounts.otp.send_otp"):
            self.client.post(reverse("accounts:login"), {"phone": "09121234567"}, HTTP_USER_AGENT=UA_DESK)
        self.assertEqual(list(Event.objects.values_list("name", flat=True)), [Event.OTP_REQUESTED])

        code = "123456"
        with mock.patch("accounts.otp.verify_code", return_value=(True, "")):
            self.client.post(reverse("accounts:verify"),
                             {"code": code, "full_name": "علی", "organization": "دانشگاه", "accept_terms": "on"},
                             HTTP_USER_AGENT=UA_DESK)
        user = User.objects.get(phone="09121234567")
        v = Visitor.objects.get()
        self.assertEqual(v.user, user)
        self.assertEqual((v.first_source, v.first_campaign), ("linkedin", "oct"))
        names = set(Event.objects.values_list("name", flat=True))
        self.assertTrue({Event.OTP_REQUESTED, Event.OTP_VERIFIED, Event.SIGNUP} <= names)

        rows = reports.breakdown(["first_source", "first_medium", "first_campaign"], 30)
        self.assertEqual(rows[0]["label"], "linkedin / paid / oct")
        self.assertEqual((rows[0]["visitors"], rows[0]["requested"], rows[0]["signup"], rows[0]["stuck"]), (1, 1, 1, 0))
        self.assertEqual(reports.funnel(30)[3]["n"], 1)

    def test_requested_but_never_verified_is_stuck(self):
        self.client.get("/?utm_source=telegram&utm_medium=social", HTTP_USER_AGENT=UA_DESK)
        with mock.patch("accounts.otp.send_otp"):
            self.client.post(reverse("accounts:login"), {"phone": "09121234567"}, HTTP_USER_AGENT=UA_DESK)
        row = reports.breakdown(["first_source"], 30)[0]
        self.assertEqual((row["requested"], row["signup"], row["stuck"]), (1, 0, 1))

    def test_funnel_counts_import_and_map_steps_once_per_visitor(self):
        self.client.get("/", HTTP_USER_AGENT=UA_DESK)
        req = SimpleNamespace(COOKIES={"gp_vid": Visitor.objects.get().vid}, META={}, path="/dashboard/1/")
        for name in (Event.IMPORT_QUEUED, Event.MAP_VIEWED, Event.MAP_VIEWED):
            services.record(req, name)
        steps = {s["label"]: s["n"] for s in reports.funnel(30)}
        self.assertEqual((steps["ورود اولین داده"], steps["دیدن داده روی نقشه"]), (1, 1))

    def test_no_cookie_no_event(self):
        services.record(SimpleNamespace(COOKIES={}, META={}, path="/x"), "signup")
        self.assertEqual(Event.objects.count(), 0)


class ReportAndCleanupTests(TestCase):
    def test_admin_report_renders_for_superuser(self):
        admin = User.objects.create_admin("root", "Str0ng-Admin-Pass")
        self.client.force_login(admin)
        self.client.logout()
        self.client.force_login(admin)
        Visitor.objects.create(vid="a" * 32, first_source="linkedin", first_medium="paid", first_landing="/")
        r = self.client.get(reverse("admin:analytics_sourcereport_changelist") + "?days=7")
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "linkedin / paid")
        self.assertContains(r, "قیف ثبت‌نام")

    def test_cleanup_keeps_linked_visitors(self):
        user = User.objects.create_user("09120000000", full_name="x")
        old = timezone.now() - timedelta(days=200)
        keep = Visitor.objects.create(vid="b" * 32, user=user)
        drop = Visitor.objects.create(vid="c" * 32)
        Visitor.objects.update(first_seen=old, last_seen=old)
        Event.objects.create(visitor=drop, name="signup")
        Event.objects.update(created_at=old)
        res = services.cleanup(90)
        self.assertEqual(res["visitors"], 1)
        self.assertTrue(Visitor.objects.filter(pk=keep.pk).exists())
        self.assertFalse(Visitor.objects.filter(pk=drop.pk).exists())


class TrackedLinkTests(TestCase):
    def setUp(self):
        from analytics.models import TrackedLink

        self.link = TrackedLink.objects.create(
            slug="linkedin-oct", name="لینکدین مهر", destination="/?utm_source=old&x=1", utm_source="linkedin",
            utm_medium="paid", utm_campaign="oct-gis", utm_content="img1")

    def test_target_path_replaces_old_utm_and_keeps_other_params(self):
        t = self.link.target_path()
        self.assertTrue(t.startswith("/?x=1&utm_source=linkedin&utm_medium=paid&utm_campaign=oct-gis"))
        self.assertNotIn("old", t)

    def test_redirect_counts_click_and_attribution_flows_end_to_end(self):
        r = self.client.get("/go/linkedin-oct/", HTTP_USER_AGENT=UA_DESK)
        self.assertEqual(r.status_code, 302)
        self.assertEqual(r["Cache-Control"], "no-store")
        self.link.refresh_from_db()
        self.assertEqual(self.link.clicks, 1)
        self.client.get(r["Location"], HTTP_USER_AGENT=UA_DESK)
        v = Visitor.objects.get()
        self.assertEqual((v.first_source, v.first_medium, v.first_campaign), ("linkedin", "paid", "oct-gis"))
        self.assertEqual(Visit.objects.get().content, "img1")

    def test_bot_preview_not_counted_but_still_redirects(self):
        r = self.client.get("/go/linkedin-oct/", HTTP_USER_AGENT="TelegramBot (like TwitterBot)")
        self.assertEqual(r.status_code, 302)
        self.link.refresh_from_db()
        self.assertEqual(self.link.clicks, 0)

    def test_inactive_or_unknown_is_404(self):
        self.link.is_active = False
        self.link.save()
        self.assertEqual(self.client.get("/go/linkedin-oct/").status_code, 404)
        self.assertEqual(self.client.get("/go/nope/").status_code, 404)

    def test_admin_validation_and_normalization(self):
        from analytics.admin import TrackedLinkForm

        base = {"slug": "Bale-Ad", "name": "بله", "destination": "/", "utm_source": " Bale ", "utm_medium": "Social",
                "utm_campaign": "Mehr Ad", "is_active": True}
        form = TrackedLinkForm(base)
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual((form.cleaned_data["slug"], form.cleaned_data["utm_source"], form.cleaned_data["utm_campaign"]),
                         ("bale-ad", "bale", "mehr-ad"))
        for bad in ("https://evil.com", "//evil.com", "evil.com/x", "/\\evil.com"):
            self.assertFalse(TrackedLinkForm({**base, "slug": "x", "destination": bad}).is_valid(), bad)

    def test_admin_pages_render(self):
        admin = User.objects.create_admin("root", "Str0ng-Admin-Pass")
        self.client.force_login(admin)
        r = self.client.get(reverse("admin:analytics_trackedlink_changelist"))
        self.assertContains(r, "linkedin-oct")
        r = self.client.get(reverse("admin:analytics_trackedlink_change", args=[self.link.pk]))
        self.assertContains(r, "/go/linkedin-oct/")
        self.assertEqual(self.client.get(reverse("admin:analytics_trackedlink_add")).status_code, 200)


class ClassifyRepairTests(SimpleTestCase):
    def test_broken_utm_with_question_marks_is_split(self):
        r = cl.classify({"utm_source": "aparat?utm_medium=social?utm_campaign=help-videos"}, "")
        self.assertEqual((r["source"], r["medium"], r["campaign"]), ("aparat", "social", "help-videos"))

    def test_proper_params_not_overwritten_by_repair(self):
        r = cl.classify({"utm_source": "x?utm_medium=y", "utm_medium": "paid"}, "")
        self.assertEqual((r["source"], r["medium"]), ("x", "paid"))

    def test_android_app_referrer_is_social_not_referral(self):
        self.assertEqual(cl.classify({}, "com.linkedin.android")["source"], "linkedin")
        r = cl.classify({}, "org.telegram.messenger")
        self.assertEqual((r["source"], r["medium"]), ("telegram", "social"))


class ActionTrackingTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(phone="09121234567")
        self.client.get("/", HTTP_USER_AGENT=UA_DESK)
        self.client.force_login(self.user)

    def test_view_and_post_are_recorded_without_form_content(self):
        self.client.get(reverse("databases:create"), HTTP_USER_AGENT=UA_DESK)
        self.client.post(reverse("databases:create"), {"label": "x", "password": "secret-pass-123"}, HTTP_USER_AGENT=UA_DESK)
        names = set(Event.objects.values_list("name", flat=True))
        self.assertTrue(any(n.startswith("view:") for n in names), names)
        self.assertTrue(any(n.startswith("do:") for n in names), names)
        blob = " ".join(str(e.meta) + e.path for e in Event.objects.all())
        self.assertNotIn("secret-pass-123", blob)

    def test_staff_and_ajax_are_not_recorded(self):
        staff = User.objects.create_admin("root", "Str0ng-Admin-Pass")
        self.client.force_login(staff)
        self.client.get(reverse("databases:create"), HTTP_USER_AGENT=UA_DESK)
        self.client.force_login(self.user)
        self.client.get(reverse("databases:create"), HTTP_USER_AGENT=UA_DESK, HTTP_X_REQUESTED_WITH="fetch")
        self.assertFalse(Event.objects.filter(name__startswith="view:").exists())


class ExportTests(TestCase):
    def test_workbook_has_all_sheets_and_masks_phone(self):
        import io

        from openpyxl import load_workbook

        from analytics import export

        user = User.objects.create_user(phone="09121234567")
        self.client.get("/?utm_source=linkedin&utm_medium=paid&utm_campaign=oct", HTTP_USER_AGENT=UA_DESK)
        self.client.force_login(user)
        self.client.get(reverse("databases:create"), HTTP_USER_AGENT=UA_DESK)
        wb = load_workbook(io.BytesIO(export.build(30)))
        for name in ("راهنما", "خلاصه", "قیف", "روزانه", "منابع", "کاربران", "ریزکارها", "رویدادهای قیف", "بازدیدها",
                     "ثبت عملیات", "پایگاه‌ها", "کارها", "فاکتورها", "لینک‌ها"):
            self.assertIn(name, wb.sheetnames)
        users = [[c.value for c in row] for row in wb["کاربران"].iter_rows(min_row=2)]
        self.assertEqual(users[0][1], "0912***4567")
        self.assertEqual(users[0][-1], "فقط ثبت‌نام")
        actions = [[c.value for c in row] for row in wb["ریزکارها"].iter_rows(min_row=2)]
        self.assertTrue(any(r[3] == "databases:create" for r in actions))
        self.assertEqual(export.mask_phone("تماس 09121234567 و 09351112233"), "تماس 0912***4567 و 0935***2233")

    def test_admin_download_requires_staff_and_returns_xlsx(self):
        self.assertEqual(self.client.get(reverse("admin:analytics_export")).status_code, 302)
        admin = User.objects.create_admin("root", "Str0ng-Admin-Pass")
        self.client.force_login(admin)
        r = self.client.get(reverse("admin:analytics_export") + "?days=0")
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheetml", r["Content-Type"])
        self.assertTrue(r.content.startswith(b"PK"))
