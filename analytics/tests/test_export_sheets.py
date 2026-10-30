import io

from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook

from accounts.models import User
from analytics import export


class ExportSheetsTests(TestCase):
    """هر برگه با داده‌ی واقعی ساخته می‌شود تا نام فیلد اشتباه در ستون‌ها (که با جدول خالی پیدا نمی‌شود) لو برود."""

    def test_all_builders_with_rows(self):
        from analysis.models import Question, ToolRun
        from assistant.models import Usage
        from billing.models import Invoice
        from databases.models import ActivityLog, Plan, UserDatabase
        from jobs.models import Job
        from kb.models import Article, Category, Feedback, MissedQuery
        from maps.models import WebMap
        from publish.models import ApiKey, Layer, Service
        from support.models import Ticket
        from survey.models import SurveyProject

        u = User.objects.create_user(phone="09121234567")
        t = User.objects.create_user(phone="09351112233")
        plan = Plan.objects.first()
        from databases.models import DatabaseServer
        srv = DatabaseServer.objects.create(name="s", host="h")
        db = UserDatabase.objects.create(server=srv, owner=u, plan=plan, label="a", db_name="a", role_name="a_user")
        ActivityLog.objects.create(user=u, database=db, action="team_invite", message="دعوت 09121234567")
        Job.objects.create(kind="import", status="done", database=db, user=u)
        Invoice.objects.create(number="F1", user=u, kind="new", plan=plan, amount_toman=1000, status="paid", expires_at=timezone.now(), paid_at=timezone.now())
        Ticket.objects.create(number="T1", user=u, subject="s")
        WebMap.objects.create(database=db, slug="m", title="m", created_by=u)
        svc = Service.objects.create(database=db)
        Layer.objects.create(service=svc, table="t", geom_column="g", srid=4326, geom_type="Point")
        ApiKey.objects.create(database=db, created_by=u, name="k", prefix="abc", key_hash="x")
        SurveyProject.objects.create(database=db, slug="s", name="s", srid=32639, created_by=u)
        Usage.objects.create(user=u, actor="u", status="ok", mode="rag", path="/")
        MissedQuery.objects.create(query_norm="q", sample="q")
        cat = Category.objects.create(slug="c", title="c")
        art = Article.objects.create(category=cat, slug="a", title="a")
        Feedback.objects.create(article=art, helpful=True)
        ToolRun.objects.create(database=db, user=u, tool="buffer", rows=3)
        Question.objects.create(database=db, user=u, text="سلام")
        wb = load_workbook(io.BytesIO(export.build(0)))
        rows = {ws.title: ws.max_row for ws in wb.worksheets}
        for name in ("کاربران", "پایگاه‌ها", "کارها", "تحلیل‌ها", "نقشه‌ها", "لایه‌های منتشرشده", "کلیدهای API",
                     "فاکتورها", "پشتیبانی", "برداشت", "دستیار", "جست‌وجوی بی‌نتیجه", "بازخورد مقاله", "ثبت عملیات"):
            self.assertGreaterEqual(rows[name], 2, name)
        log = [c.value for c in list(wb["ثبت عملیات"].iter_rows(min_row=2, max_row=2))[0]]
        self.assertEqual(log[-1], "دعوت 0912***4567")
        keys = {c.value for ws in wb.worksheets for row in ws.iter_rows() for c in row if isinstance(c.value, str)}
        self.assertFalse(any("secret" in k or "key_hash" in k for k in keys))
        self.assertEqual(t.pk is not None, True)
